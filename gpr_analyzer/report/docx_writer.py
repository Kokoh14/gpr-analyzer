"""Word (.docx) renderer (python-docx).

Word cannot embed fonts reliably, so Thai text uses a complex-script font that ships with
Windows ("Leelawadee UI" by default; "TH Sarabun New" is the Thai government standard).
"""

from __future__ import annotations

import io
import os
from pathlib import Path

from gpr_analyzer.report.compose import Bullets, Figure, Heading, PageBreak, Para, Report, Table

ACTION_FILL = {"E": "F8D7DA", "1": "FDE2CC", "2": "FFF3CD", "3": "D6E6FA", "none": "FFFFFF"}
HEAD_FILL = "E8ECF1"
MUTED = (0x5F, 0x66, 0x70)


def _set_fonts(target, latin: str, thai: str, size_pt: float | None = None) -> None:
    """Set Latin + complex-script (Thai) font on a style or run, including the cs size."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    font = target.font
    font.name = latin
    if size_pt:
        font.size = Pt(size_pt)
    rpr = target.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attr), latin)
    rfonts.set(qn("w:cs"), thai)
    if size_pt:
        szcs = rpr.find(qn("w:szCs"))
        if szcs is None:
            szcs = OxmlElement("w:szCs")
            rpr.append(szcs)
        szcs.set(qn("w:val"), str(int(round(size_pt * 2))))


def _shade(cell, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def render_docx(report: Report, path: str | os.PathLike | None = None,
                latin_font: str = "Calibri", thai_font: str = "Leelawadee UI") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.8))
    usable = sec.page_width - sec.left_margin - sec.right_margin

    _set_fonts(doc.styles["Normal"], latin_font, thai_font, 10)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        _set_fonts(doc.styles[name], latin_font, thai_font, size)
    _set_fonts(doc.styles["List Bullet"], latin_font, thai_font, 10)

    footer = sec.footer.paragraphs[0]
    footer.text = report.footer
    footer.runs[0].font.size = Pt(8)
    if report.draft:
        hdr = sec.header.paragraphs[0]
        run = hdr.add_run(" · ".join(report.draft_text))
        run.font.size, run.font.bold = Pt(8), True
        run.font.color.rgb = RGBColor(0xB0, 0x00, 0x2A)
        hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(report.title, style="Title")
    p = doc.add_paragraph(report.subtitle)
    p.runs[0].font.color.rgb = RGBColor(*MUTED)
    if report.draft:
        p = doc.add_paragraph()
        r = p.add_run("\n".join(report.draft_text))
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xB0, 0x00, 0x2A)

    for b in report.blocks:
        if isinstance(b, Heading):
            doc.add_heading(b.text, level=min(b.level, 3))
        elif isinstance(b, Para):
            p = doc.add_paragraph()
            r = p.add_run(b.text)
            if b.style in ("secondary", "small", "caption"):
                r.font.color.rgb = RGBColor(*MUTED)
                r.font.size = Pt(9)
                if b.style == "secondary":
                    r.font.italic = True
        elif isinstance(b, Bullets):
            for item in b.items:
                doc.add_paragraph(item, style="List Bullet")
        elif isinstance(b, Table):
            ncols = len(b.header or b.rows[0])
            t = doc.add_table(rows=0, cols=ncols)
            t.style = "Table Grid"
            widths = b.widths or [1] * ncols
            tot = sum(widths)
            size = Pt(8 if ncols > 6 else 9)
            rows = ([b.header] if b.header else []) + b.rows
            for i, r in enumerate(rows):
                cells = t.add_row().cells
                is_head = bool(b.header) and i == 0
                tag = None if is_head or not b.row_tags else b.row_tags[i - (1 if b.header else 0)]
                for j, txt in enumerate(r):
                    cells[j].width = int(usable * widths[j] / tot)
                    para = cells[j].paragraphs[0]
                    run = para.add_run(str(txt))
                    run.font.size = size
                    if is_head:
                        run.font.bold = True
                        _shade(cells[j], HEAD_FILL)
                    elif b.kind == "kv" and j == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(*MUTED)
                    elif tag is not None and j == 0:
                        run.font.bold = True
                        _shade(cells[j], ACTION_FILL.get(tag, "FFFFFF"))
            doc.add_paragraph()
        elif isinstance(b, Figure):
            w = int(usable * b.width)
            max_h = (sec.page_height - sec.top_margin - sec.bottom_margin) * 0.7
            if w * b.aspect > max_h:
                w = int(max_h / b.aspect)
            doc.add_picture(io.BytesIO(b.png), width=w)
            for line in b.caption:
                p = doc.add_paragraph()
                r = p.add_run(line)
                r.font.size, r.font.italic = Pt(8.5), True
                r.font.color.rgb = RGBColor(*MUTED)
        elif isinstance(b, PageBreak):
            doc.add_page_break()

    buf = io.BytesIO()
    doc.save(buf)
    out = buf.getvalue()
    if path is not None:
        Path(path).write_bytes(out)
    return out
